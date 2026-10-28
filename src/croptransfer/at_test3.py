"""Third Austria test: Styria with the Hatzendorf Excel series (protocol `reports/freeze_protocol_at3.md`, tag freeze-at-3).

New compared with freeze-at-2: the yields come from the long N-fertilisation series of Wagna and Kalsdorf
(Excel files of the Hatzendorf station, received 2026-10-28) instead of the hand-transcribed report rows,
and 2026 is added. Scores 2018-2025 are taken unchanged from the first test (`site_year.csv`); 2026 scores
are computed later with the frozen rules (same parcel sampling, HLS extraction and stress score).

Usage:  python -m croptransfer.at_test3 --yields-only   (tidy yield table + site-year yields, no scores read)
        python -m croptransfer.at_test3                 (full evaluation, only after freeze-at-3)
"""

from __future__ import annotations

import argparse
import json
import re

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from croptransfer import at_test
from croptransfer.at_test_stmk import evaluate
from croptransfer.config import load_config, repo_path

YEAR_HEADER = re.compile(r"^Ertrag (\d{4})(?!.*t/ha)")


def read_excel_yields(path, sheet: str, header_row: int, variant_col: int) -> pd.DataFrame:
    """Variant x year grain yields (dt/ha) of one 'Daten Ertrag' sheet -> site-less long table.

    Year columns are the header cells 'Ertrag YYYY ...' (t/ha duplicates skipped); the header text after
    the year is kept as `header_note` (e.g. 'nicht mitgewertet'). Empty or non-numeric cells are dropped.
    """
    ws = load_workbook(path, data_only=True, read_only=True)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[header_row - 1]
    years = {}
    for j, h in enumerate(header):
        m = YEAR_HEADER.match(str(h or "").strip())
        if m and int(m.group(1)) not in years.values():
            years[j] = int(m.group(1))
    out, seen = [], set()
    for r in rows[header_row:]:
        var = r[variant_col - 1]
        if var is None or str(var).strip() == "" or str(var).strip() in seen:
            continue                                    # later blocks of the sheet reuse codes (sums, helper tables)
        seen.add(str(var).strip())
        for j, yr in years.items():
            v = r[j]
            if isinstance(v, (int, float)) and not isinstance(v, bool) and v > 0:
                note = str(header[j]).split(str(yr), 1)[1].strip()
                out.append({"variant": str(var).strip(), "year": yr, "yield_dt_ha": float(v), "header_note": note})
    return pd.DataFrame(out)


def tidy_yields(cfg) -> pd.DataFrame:
    """All configured Excel files -> site, year, variant, yield_dt_ha, header_note, excluded."""
    t3 = cfg["at_test3"]
    parts = []
    for site, s in t3["excel"].items():
        d = read_excel_yields(repo_path(s["path"]), s["sheet"], s["header_row"], s["variant_col"])
        d.insert(0, "site", site)
        d["excluded"] = d.year.isin(s.get("exclude_years", []))
        parts.append(d)
    return pd.concat(parts, ignore_index=True)


def site_year_yields(tidy: pd.DataFrame, first_year: int, core_only: bool = True) -> pd.DataFrame:
    """Y(c,y) in t/ha: mean over the core variants (yield in every counted year >= first_year), else all variants."""
    d = tidy[~tidy.excluded & (tidy.year >= first_year)]
    out = []
    for site, g in d.groupby("site"):
        years = g.year.nunique()
        core = g.groupby("variant").year.nunique()
        use = g[g.variant.isin(core[core == years].index)] if core_only else g
        m = use.groupby("year").yield_dt_ha.agg(["mean", "size"]).reset_index()
        out += [{"site": site, "year": int(r.year), "y_t_ha": r["mean"] / 10.0, "n_variants": int(r["size"])}
                for _, r in m.iterrows()]
    return pd.DataFrame(out)


def rows(cfg, scores: pd.DataFrame, core_only: bool = True) -> pd.DataFrame:
    """Frozen site-year scores merged with YAS_site from the Excel series (rule of `at_test.yas_site`, unchanged)."""
    t3 = cfg["at_test3"]
    y = at_test.yas_site(site_year_yields(tidy_yields(cfg), t3["first_yield_year"], core_only)[["site", "year", "y_t_ha"]])
    base = scores[scores.site.isin(t3["sites"]) & scores.year.between(*t3["years"])]
    return base.merge(y, on=["site", "year"], how="left")


def load_scores(cfg) -> pd.DataFrame:
    """Scores 2018-2025 from the first test plus the 2026 scores file (computed after the freeze)."""
    t3 = cfg["at_test3"]
    keep = ["site", "year", "ndvi_base"] + [f"{p}_{c}" for c in cfg["at_test"]["cutoffs"] for p in ("score", "pred", "n_valid")]
    first = pd.read_csv(repo_path(cfg["at_test_stmk"]["first_test_results"]))
    new = pd.read_csv(repo_path(t3["scores_2026"]))
    return pd.concat([first[[c for c in keep if c in first]], new[[c for c in keep if c in new]]], ignore_index=True)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--yields-only", action="store_true")
    a = ap.parse_args()
    cfg = load_config()
    t3 = cfg["at_test3"]
    if a.yields_only:
        tidy = tidy_yields(cfg)
        tidy.to_csv(repo_path(t3["tidy_csv"]), index=False)
        print(site_year_yields(tidy, t3["first_yield_year"]).to_string(index=False))
        return
    out = repo_path(t3["results_dir"])
    out.mkdir(parents=True, exist_ok=True)
    scores = load_scores(cfg)
    res = {"protocol": "reports/freeze_protocol_at3.md", "tag": "freeze-at-3"}
    for name, core in [("primary_core_variants", True), ("sensitivity_all_variants", False)]:
        df = rows(cfg, scores, core)
        if core:
            df.to_csv(out / "site_year.csv", index=False)
        res[name] = evaluate(df, cfg)       # same metrics and thresholds as freeze-at-2 (min 15 site-years)
        if core:
            seen = (df.site + "|" + df.year.astype(str)).isin(t3["seen_new_points"])
            res["sensitivity_without_seen_new_points"] = evaluate(df[~seen], cfg)
            res["blind_site_years_2026"] = df[df.year == 2026].to_dict("records")
    (out / "metrics.json").write_text(json.dumps(res, indent=2, default=float, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(res, indent=1, default=float, ensure_ascii=False)[:6000])


if __name__ == "__main__":
    main()
