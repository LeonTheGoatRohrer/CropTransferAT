"""Third Styria test: Excel parsing and site-year yields on a synthetic workbook (no real yield is read)."""

import pandas as pd
from openpyxl import Workbook

from croptransfer import at_test3
from croptransfer.config import load_config


def _workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Daten"
    ws.cell(5, 2, "Variante")
    for j, h in enumerate(["Ertrag 2018", "Ertrag 2019 (Feldtag - nicht mitgewertet)", "Ertrag 2020", "Ertrag 2020 t/ha"], start=3):
        ws.cell(5, j, h)
    rows = [("0", [50, 40, 60, 6.0]), ("A", [100, None, 120, 12.0]), ("Z1", [None, None, 130, 13.0]), (None, [999, 999, 999, 99]),
            ("0", [6451, 6451, 6451, 99])]                  # helper block further down reusing a code
    for i, (var, vals) in enumerate(rows, start=6):
        ws.cell(i, 2, var)
        for j, v in enumerate(vals, start=3):
            ws.cell(i, j, v)
    wb.save(path)


def test_read_excel_yields_skips_t_ha_columns_empty_cells_and_rows_without_variant(tmp_path):
    p = tmp_path / "w.xlsx"
    _workbook(p)
    d = at_test3.read_excel_yields(p, "Daten", header_row=5, variant_col=2)
    assert set(d.year) == {2018, 2019, 2020}
    assert len(d) == 6                                      # 0: 3 years, A: 2, Z1: 1; rows without or with a repeated code ignored
    assert d.yield_dt_ha.max() == 130                       # t/ha duplicate column never read
    assert d.loc[d.year == 2019, "header_note"].iloc[0] == "(Feldtag - nicht mitgewertet)"


def test_site_year_yields_core_variants_and_exclusions():
    tidy = pd.DataFrame({
        "site": "x", "variant": ["0", "A", "Z1", "0", "A", "0", "A"],
        "year": [2020, 2020, 2020, 2021, 2021, 2012, 2012],
        "yield_dt_ha": [60, 120, 130, 50, 100, 10, 10], "excluded": False})
    core = at_test3.site_year_yields(tidy, first_year=2013)
    assert core.set_index("year").y_t_ha.to_dict() == {2020: 9.0, 2021: 7.5}     # Z1 is not in every year -> left out
    allv = at_test3.site_year_yields(tidy, first_year=2013, core_only=False)
    assert abs(allv.set_index("year").y_t_ha[2020] - 31 / 3) < 1e-9
    excl = at_test3.site_year_yields(tidy.assign(excluded=tidy.year == 2021), first_year=2013)
    assert excl.year.tolist() == [2020] and excl.n_variants.tolist() == [3]       # 2021 excluded -> Z1 becomes core


def test_config_seen_points_inside_test_window():
    t3 = load_config()["at_test3"]
    for key in t3["seen_new_points"]:
        site, year = key.split("|")
        assert site in t3["sites"] and t3["years"][0] <= int(year) <= t3["years"][1]


def test_frozen_config_and_hashes_of_third_protocol():
    import json
    import re

    from croptransfer import regional_cybench as rc
    from croptransfer.config import repo_path
    frozen = json.loads(repo_path("reports/freeze_at/at_test3_config.json").read_text(encoding="utf-8"))
    assert load_config()["at_test3"] == frozen
    text = repo_path("reports/freeze_protocol_at3.md").read_text(encoding="utf-8")
    rows = re.findall(r"\|\s*`((?:reports|results)/[^`]+)`\s*\|\s*`([0-9a-f]{64})`\s*\|", text)
    assert len(rows) == 2
    for rel, digest in rows:
        assert rc.sha256_file(repo_path(rel)) == digest, rel
